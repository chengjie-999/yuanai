"""分析结果导出 — Excel (openpyxl) + HTML (Jinja2)"""
import os
import io
import base64
import json as _json
from utils.data_path import root_path

EXPORT_DIR = os.path.join(root_path(), "data", "analysis")


def export_excel(result: dict) -> bytes:
    """Generate Excel report with stats + correlation + chart images."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="Microsoft YaHei", size=9)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def _write_header(ws, row, headers):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

    # Sheet 1: Stats + Missing
    ws1 = wb.active
    ws1.title = "统计"
    ws1.append([f"数据集: {result.get('name', '')}  |  {result.get('row_count', '')} 行"])

    describe = result.get("describe", {})
    if describe:
        ws1.append([])
        ws1.append(["数值列统计"])
        stat_keys = list(next(iter(describe.values()), {}).keys())
        _write_header(ws1, ws1.max_row + 1, ["列名"] + stat_keys)
        for col, stats in describe.items():
            row_data = [col] + [stats.get(k, "") for k in stat_keys]
            r = ws1.max_row + 1
            for c, v in enumerate(row_data, 1):
                cell = ws1.cell(row=r, column=c, value=v)
                cell.font = cell_font
                cell.border = thin_border

    missing = result.get("missing", {})
    if missing:
        ws1.append([])
        ws1.append(["缺失值"])
        for col, cnt in missing.items():
            ws1.append([col, cnt])

    # Sheet 2: Correlation
    corr = result.get("corr", [])
    labels = result.get("corr_labels", [])
    if corr and labels:
        ws2 = wb.create_sheet("相关系数")
        _write_header(ws2, 1, [""] + labels)
        for i, (label, row) in enumerate(zip(labels, corr)):
            r = i + 2
            ws2.cell(row=r, column=1, value=label).font = Font(name="Microsoft YaHei", size=9, bold=True)
            for j, v in enumerate(row):
                cell = ws2.cell(row=r, column=j + 2, value=round(v, 2))
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

    # Sheet 3+: Charts (embedded PNG images)
    charts = result.get("charts", [])
    for ch in charts:
        name = ch.get("name", "chart")[:31]
        ws = wb.create_sheet(name)
        data = ch.get("data", "")
        if data:
            img_bytes = base64.b64decode(data)
            img_io = io.BytesIO(img_bytes)
            from openpyxl.drawing.image import Image as XlImage
            xl_img = XlImage(img_io)
            xl_img.width, xl_img.height = 600, 300
            ws.add_image(xl_img, "A1")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def export_html(result: dict) -> str:
    """Generate self-contained HTML report with Jinja2."""
    import jinja2

    tmpl_dir = os.path.dirname(os.path.abspath(__file__))
    env = jinja2.Environment(loader=jinja2.FileSystemLoader(tmpl_dir), autoescape=True)

    describe_rows = []
    describe = result.get("describe", {})
    stat_keys = list(next(iter(describe.values()), {}).keys())
    for col, stats in describe.items():
        describe_rows.append({"name": col, "values": [stats.get(k, "") for k in stat_keys]})

    corr = result.get("corr", [])
    labels = result.get("corr_labels", [])
    corr_rows = []
    for i, (label, row) in enumerate(zip(labels, corr)):
        corr_rows.append({"name": label, "values": [round(v, 2) for v in row]})

    charts = []
    for ch in result.get("charts", []):
        charts.append({
            "name": ch.get("name", ""),
            "html_url": ch.get("html_url", ""),
            "img_src": f"data:image/png;base64,{ch['data']}" if ch.get("data") else "",
        })

    ts = result.get("time_series", {})

    try:
        tmpl = env.get_template("report_template.html")
        return tmpl.render(
            name=result.get("name", ""),
            row_count=result.get("row_count", 0),
            col_count=len(result.get("columns", [])),
            num_count=len(result.get("num_cols", [])),
            stat_keys=stat_keys,
            describe_rows=describe_rows,
            missing=result.get("missing", {}),
            corr_labels=labels,
            corr_rows=corr_rows,
            charts=charts,
            time_series=ts,
        )
    except Exception:
        return _fallback_html(result, stat_keys, describe_rows, labels, corr_rows, charts, ts)


def _fallback_html(result, stat_keys, describe_rows, labels, corr_rows, charts, ts):
    """Inline HTML fallback if template not found."""
    body = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>{result['name']}</title>
<style>body{{font-family:"Microsoft YaHei",sans-serif;max-width:900px;margin:0 auto;padding:20px;color:#333}}
table{{border-collapse:collapse;width:100%;margin:8px 0}}th,td{{border:1px solid #ddd;padding:4px 8px;font-size:12px}}
th{{background:#4472C4;color:#fff}}h1{{font-size:18px}}h2{{font-size:14px;margin-top:20px}}
.chart{{margin:10px 0;border:1px solid #eee;border-radius:6px;overflow:hidden}}
.chart-title{{background:#f5f5f5;padding:6px 12px;font-size:12px;color:#666}}</style></head><body>
<h1>{result['name']}</h1><p>{result['row_count']} 行</p>"""

    if describe_rows:
        body += "<h2>数值列统计</h2><table><tr><th>列名</th>"
        for k in stat_keys:
            body += f"<th>{k}</th>"
        body += "</tr>"
        for row in describe_rows:
            body += f"<tr><td><b>{row['name']}</b></td>" + "".join(f"<td>{v}</td>" for v in row["values"]) + "</tr>"
        body += "</table>"

    miss = result.get("missing", {})
    if miss:
        body += "<h2>缺失值</h2><ul>" + "".join(f"<li>{k}: {v}</li>" for k, v in miss.items()) + "</ul>"

    if corr_rows:
        body += "<h2>相关系数矩阵</h2><table><tr><th></th>" + "".join(f"<th>{l}</th>" for l in labels) + "</tr>"
        for row in corr_rows:
            body += f"<tr><td><b>{row['name']}</b></td>" + "".join(f"<td>{v}</td>" for v in row["values"]) + "</tr>"
        body += "</table>"

    ts_data = result.get("time_series", {})
    if ts_data.get("anomalies"):
        body += "<h2>异常点</h2><ul>"
        for a in ts_data["anomalies"]:
            body += f"<li>{a['date']}: {a['value']} (z={a['z_score']})</li>"
        body += "</ul>"

    for ch in charts:
        body += f"<div class='chart'><div class='chart-title'>{ch['name']}</div>"
        if ch.get("img_src"):
            body += f"<img src='{ch['img_src']}' style='width:100%'>"
        body += "</div>"

    body += "</body></html>"
    return body
